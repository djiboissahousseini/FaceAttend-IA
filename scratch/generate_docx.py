
from docx import Document
from docx.shared import Inches
import os

def create_report_supplement():
    doc = Document()
    doc.add_heading('FaceAttend - Annexes Techniques (Diagrammes)', 0)

    # 1. Diagrammes de Cas d'Utilisation
    doc.add_heading('I. Diagrammes de Cas d\'Utilisation', level=1)
    
    uc_dir = "/home/usain/Bureau/FaceAttend/docs/diagrams/cas_utilisation"
    uc_files = [
        ("uc_admin_utilisateurs.png", "Cas d'utilisation : Administration des Utilisateurs (Auth + CRUD)"),
        ("uc_admin_entites_operations.png", "Cas d'utilisation : Opérations Critiques du Système (IA, Séances, Rapports)"),
        ("uc_enseignant_final.png", "Cas d'utilisation : Espace Enseignant (Suivi et Terminal Caméra)"),
        ("uc_etudiant_final.png", "Cas d'utilisation : Espace Étudiant (Portail et Mobile)"),
        ("uc_camera_ia.png", "Cas d'utilisation : Système de Reconnaissance Faciale (Cœur IA)")
    ]

    for filename, caption in uc_files:
        path = os.path.join(uc_dir, filename)
        if os.path.exists(path):
            doc.add_heading(caption, level=2)
            doc.add_picture(path, width=Inches(6))
            doc.add_paragraph(f"Figure : {caption}")

    # 2. Diagrammes de Séquence
    doc.add_heading('II. Diagrammes de Séquence Détaillés', level=1)
    
    seq_dir = "/home/usain/Bureau/FaceAttend/docs/diagrams/sequences_detaillees"
    seq_files = [
        ("seq_ajouter.png", "Séquence : Ajout d'un nouvel utilisateur (Enregistrement Empreinte)"),
        ("seq_modifier.png", "Séquence : Modification des données utilisateur"),
        ("seq_rechercher.png", "Séquence : Recherche multicritère"),
        ("seq_supprimer.png", "Séquence : Suppression sécurisée"),
        ("seq_creer_seance.png", "Séquence : Initialisation et Ouverture d'une séance d'émargement")
    ]

    for filename, caption in seq_files:
        path = os.path.join(seq_dir, filename)
        if os.path.exists(path):
            doc.add_heading(caption, level=2)
            doc.add_picture(path, width=Inches(6))
            doc.add_paragraph(f"Figure : {caption}")

    output_path = "/home/usain/Bureau/FaceAttend/docs/FaceAttend_Diagrammes_Rapport.docx"
    doc.save(output_path)
    print(f"Document généré avec succès : {output_path}")

if __name__ == "__main__":
    create_report_supplement()
