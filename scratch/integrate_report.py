
import os
from docx import Document
from docx.shared import Inches

def integrate_diagrams(report_path):
    print(f"Tentative d'intégration dans : {report_path}")
    
    if not os.path.exists(report_path):
        print(f"Erreur : Le fichier {report_path} n'existe pas.")
        return

    try:
        doc = Document(report_path)
        
        doc.add_page_break()
        doc.add_heading('III. Conception et Modélisation (UML)', level=1)
        doc.add_paragraph("Cette section présente les diagrammes UML modélisant les fonctionnalités et les interactions du système FaceAttend.")

        # 1. Cas d'Utilisation
        doc.add_heading('III.1. Diagrammes de Cas d\'Utilisation', level=2)
        uc_dir = "/home/usain/Bureau/FaceAttend/docs/diagrams/cas_utilisation"
        uc_files = [
            ("uc_admin_utilisateurs.png", "Administration des Utilisateurs (Auth + CRUD)"),
            ("uc_admin_entites_operations.png", "Opérations Critiques du Système (IA, Séances, Rapports)"),
            ("uc_enseignant_final.png", "Espace Enseignant (Suivi et Terminal Caméra)"),
            ("uc_etudiant_final.png", "Espace Étudiant (Portail et Mobile)"),
            ("uc_camera_ia.png", "Système de Reconnaissance Faciale (Cœur IA)")
        ]

        for filename, caption in uc_files:
            path = os.path.join(uc_dir, filename)
            if os.path.exists(path):
                doc.add_heading(caption, level=3)
                doc.add_picture(path, width=Inches(6))
                doc.add_paragraph(f"Figure : {caption}")

        # 2. Séquence
        doc.add_page_break()
        doc.add_heading('III.2. Diagrammes de Séquence Détaillés', level=2)
        seq_dir = "/home/usain/Bureau/FaceAttend/docs/diagrams/sequences_detaillees"
        seq_files = [
            ("seq_ajouter.png", "Séquence : Ajout d'un nouvel utilisateur"),
            ("seq_modifier.png", "Séquence : Modification des données"),
            ("seq_rechercher.png", "Séquence : Recherche multicritère"),
            ("seq_supprimer.png", "Séquence : Suppression sécurisée"),
            ("seq_creer_seance.png", "Séquence : Création et Ouverture de séance")
        ]

        for filename, caption in seq_files:
            path = os.path.join(seq_dir, filename)
            if os.path.exists(path):
                doc.add_heading(caption, level=3)
                doc.add_picture(path, width=Inches(6))
                doc.add_paragraph(f"Figure : {caption}")

        doc.save(report_path)
        print("Intégration réussie !")

    except Exception as e:
        print(f"Erreur lors de l'intégration : {e}")
        print("Note : Si le fichier est ouvert dans LibreOffice, veuillez le fermer avant de réessayer.")

if __name__ == "__main__":
    path = "/home/usain/Téléchargements/PCF de Application de gestion des absences basée sur la reconnaissance faciale.docx"
    integrate_diagrams(path)
